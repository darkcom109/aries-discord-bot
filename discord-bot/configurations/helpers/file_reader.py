from io import BytesIO

from pypdf import PdfReader
from pptx import Presentation
from pptx.shapes.autoshape import Shape

def extract_text(filename: str, file_bytes: bytes) -> str:
    extension = filename.rsplit(".", 1)[-1].lower()

    if extension == "pdf":
        reader = PdfReader(BytesIO(file_bytes))
        sections = []

        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                sections.append(f"Page {page_number}:\n{text}")

        return "\n\n".join(sections)

    if extension == "pptx":
        presentation = Presentation(BytesIO(file_bytes))
        sections = []

        for slide_number, slide in enumerate(presentation.slides, start=1):
            text_parts=[]

            for shape in slide.shapes:
                if isinstance(shape, Shape) and shape.has_text_frame:
                    text = shape.text.strip()
                    if text:
                        text_parts.append(text)

            if text_parts:
                sections.append(
                    f"Slide {slide_number:\n}" + "\n".join(text_parts)
                )

        return "\n\n".join(sections)

    raise ValueError(f"Unsupported file type: .{extension}")