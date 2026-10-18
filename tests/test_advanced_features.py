import asyncio
from io import BytesIO
from pathlib import Path
import sys
import shutil
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch
from uuid import uuid4

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "discord-bot"))

import discord
from pptx import Presentation
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

import database
import library_database
from clients.source_client import select_source_text, source_response
from commands.context_command import MessageQuestionModal
from configurations.helpers.file_reader import extract_text
from configurations.helpers.reminder_helper import next_reminder_due_at
from configurations.helpers.response_helper import answer_parts, shorten_text
from discord_commands import register_commands
from library_view import LibraryView
from main import MyBot
from reminder_view import ReminderView, SnoozeModal

class StorageTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.directory = Path(__file__).resolve().parent / f".aries-test-{uuid4().hex}"
        self.directory.mkdir()
        self.addCleanup(shutil.rmtree, self.directory)
        path = (self.directory / "test.db").as_posix()
        self.engine = create_async_engine(f"sqlite+aiosqlite:///{path}")
        self.addAsyncCleanup(self.engine.dispose)
        factory = async_sessionmaker(self.engine, expire_on_commit=False)
        self.patches = [
            patch.object(database, "engine", self.engine),
            patch.object(database, "SessionFactory", factory),
            patch.object(library_database, "SessionFactory", factory),
        ]
        for item in self.patches:
            item.start()
            self.addCleanup(item.stop)
        await database.create_tables()

    async def create_delivery(self, repeat="daily"):
        reminder_id = await database.save_reminder("guild", "channel", "owner", "Drink water", 100, repeat)
        async with database.SessionFactory() as session:
            reminder = await session.get(database.Reminder, reminder_id)
        delivery_id = await database.save_reminder_delivery(reminder)
        await database.set_reminder_delivery_message(delivery_id, "123")
        return reminder_id, delivery_id

    async def test_library_is_private_and_scoped_to_server(self):
        item_id = await library_database.save_library_item("guild", "channel", "owner", "Docker fix", "docker compose up")
        self.assertEqual(len(await library_database.get_user_library("guild", "owner")), 1)
        self.assertEqual(await library_database.get_user_library("guild", "stranger"), [])
        self.assertEqual(await library_database.get_user_library("other guild", "owner"), [])
        self.assertIsNone(await library_database.get_user_library_item(item_id, "guild", "stranger"))
        self.assertFalse(await library_database.delete_user_library_item(item_id, "guild", "stranger"))
        self.assertTrue(await library_database.delete_user_library_item(item_id, "guild", "owner"))
        self.assertFalse(await library_database.delete_user_library_item(item_id, "guild", "owner"))

    async def test_library_searches_title_and_content_with_literal_wildcards(self):
        await library_database.save_library_item("guild", "channel", "owner", "Docker fix", "Use compose for 100%_literal")
        await library_database.save_library_item("guild", "channel", "owner", "Other", "100percent")
        self.assertEqual(len(await library_database.get_user_library("guild", "owner", "DOCKER compose")), 1)
        self.assertEqual(len(await library_database.get_user_library("guild", "owner", "%_")), 1)
        self.assertEqual(await library_database.get_user_library("guild", "owner", "missing"), [])

    async def test_saving_same_source_updates_only_owners_copy(self):
        url = "https://discord.com/channels/1/2/3"
        first = await library_database.save_library_item("guild", "channel", "owner", "Old", "old text", url)
        updated = await library_database.save_library_item("guild", "channel", "owner", "New", "updated text", url)
        other = await library_database.save_library_item("guild", "channel", "other", "Other", "private", url)
        self.assertEqual(first, updated)
        self.assertNotEqual(first, other)
        self.assertEqual((await library_database.get_user_library_item(first, "guild", "owner")).content, "updated text")

    async def test_library_survives_reopening_and_new_tables_preserve_reminders(self):
        reminder_id, _ = await self.create_delivery()
        item_id = await library_database.save_library_item(None, "dm", "owner", "Private", "remember this")
        await self.engine.dispose()
        await database.create_tables()
        self.assertIsNotNone(await library_database.get_user_library_item(item_id, None, "owner"))
        self.assertIsNone(await library_database.get_user_library_item(item_id, "guild", "owner"))
        async with database.SessionFactory() as session:
            self.assertIsNotNone(await session.get(database.Reminder, reminder_id))

    async def test_duplicate_snooze_clicks_create_only_one_copy(self):
        reminder_id, delivery_id = await self.create_delivery()
        with patch("database.time.time", return_value=1000):
            results = await asyncio.gather(
                database.resolve_reminder_delivery(delivery_id, "owner", 600),
                database.resolve_reminder_delivery(delivery_id, "owner", 600),
            )
        self.assertEqual(sorted(results), [False, True])
        async with database.SessionFactory() as session:
            reminders = list(await session.scalars(select(database.Reminder).order_by(database.Reminder.id)))
        self.assertEqual(len(reminders), 2)
        self.assertEqual(reminders[0].id, reminder_id)
        self.assertEqual((reminders[0].repeat, reminders[0].due_at), ("daily", 100))
        self.assertEqual((reminders[1].repeat, reminders[1].due_at), ("none", 1600))

    async def test_reminder_controls_reject_other_users_and_invalid_snooze(self):
        _, delivery_id = await self.create_delivery()
        self.assertFalse(await database.resolve_reminder_delivery(delivery_id, "stranger", 600))
        self.assertFalse((await database.get_reminder_delivery(delivery_id)).resolved)
        with self.assertRaises(ValueError):
            await database.resolve_reminder_delivery(delivery_id, "owner", -60)
        with self.assertRaises(ValueError):
            await database.resolve_reminder_delivery(delivery_id, "owner", 8 * 86400)
        self.assertTrue(await database.resolve_reminder_delivery(delivery_id, "owner"))
        self.assertEqual(await database.load_reminder_deliveries(), [])

    async def test_persistent_controls_restore_from_database(self):
        _, delivery_id = await self.create_delivery()
        bot = MyBot()
        with patch.object(bot.check_reminders, "start"), patch.object(bot.commands, "sync", new=AsyncMock()):
            await bot.setup_hook()
        views = bot.persistent_views
        self.assertEqual(len(views), 1)
        expected = [item.custom_id for item in ReminderView(delivery_id).children]
        self.assertEqual([item.custom_id for item in views[0].children], expected)
        self.assertTrue(views[0].is_persistent())
        await bot.close()

    async def test_cancellation_reports_missing_or_wrong_owner(self):
        reminder_id, _ = await self.create_delivery("none")
        self.assertFalse(await database.delete_user_reminder(reminder_id, "guild", "channel", "stranger"))
        self.assertTrue(await database.delete_user_reminder(reminder_id, "guild", "channel", "owner"))
        self.assertFalse(await database.delete_user_reminder(reminder_id, "guild", "channel", "owner"))

    async def test_delivery_loop_records_message_and_marks_one_time_reminder(self):
        reminder_id = await database.save_reminder("guild", "123", "456", "Do this", 100, "none")
        bot = MyBot()
        channel = SimpleNamespace(send=AsyncMock(return_value=SimpleNamespace(id=789)))
        with patch.object(bot, "get_partial_messageable", return_value=channel):
            await bot.check_reminders.coro(bot)
        deliveries = await database.load_reminder_deliveries()
        self.assertEqual(len(deliveries), 1)
        self.assertEqual(deliveries[0].message_id, "789")
        self.assertTrue(channel.send.call_args.kwargs["view"].is_persistent())
        async with database.SessionFactory() as session:
            self.assertTrue((await session.get(database.Reminder, reminder_id)).sent)
        await bot.close()

    async def test_failed_delivery_stays_pending_without_orphaned_controls(self):
        reminder_id = await database.save_reminder("guild", "123", "456", "Do this", 100, "none")
        bot = MyBot()
        error = discord.Forbidden(SimpleNamespace(status=403, reason="Forbidden"), "No permission")
        channel = SimpleNamespace(send=AsyncMock(side_effect=error))
        with patch.object(bot, "get_partial_messageable", return_value=channel), patch("main.print"):
            await bot.check_reminders.coro(bot)
        async with database.SessionFactory() as session:
            self.assertFalse((await session.get(database.Reminder, reminder_id)).sent)
            self.assertEqual(list(await session.scalars(select(database.ReminderDelivery))), [])
        await bot.close()

    async def test_custom_snooze_stops_old_view_and_creates_pending_copy(self):
        _, delivery_id = await self.create_delivery()
        view = ReminderView(delivery_id)
        modal = SnoozeModal(delivery_id, view)
        modal.minutes._value = "25"
        interaction = SimpleNamespace(
            user=SimpleNamespace(id="owner"),
            response=SimpleNamespace(defer=AsyncMock()),
            edit_original_response=AsyncMock(),
        )
        await modal.on_submit(interaction)
        self.assertTrue(view.is_finished())
        self.assertTrue((await database.get_reminder_delivery(delivery_id)).resolved)
        self.assertIsNone(interaction.edit_original_response.call_args.kwargs["view"])

class InteractionTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.bot = MyBot()
        register_commands(self.bot)

    async def asyncTearDown(self):
        await self.bot.close()

    async def test_all_commands_and_context_menus_register(self):
        names = {command.name for command in self.bot.commands.get_commands()}
        self.assertTrue({"ask", "notes", "save", "library", "Save to Aries", "Ask Aries about this"} <= names)
        self.assertEqual(self.bot.commands.get_command("ask").name, "ask")

    async def test_oversized_image_is_acknowledged_and_never_downloaded(self):
        interaction = SimpleNamespace(response=SimpleNamespace(defer=AsyncMock()), edit_original_response=AsyncMock())
        image = SimpleNamespace(content_type="image/png", size=9 * 1024 * 1024, read=AsyncMock())
        with patch("commands.ask_command.ollama_response", new=AsyncMock()) as request:
            await self.bot.commands.get_command("ask").callback(interaction, "Look at this", image)
        interaction.response.defer.assert_awaited_once()
        image.read.assert_not_awaited()
        request.assert_not_awaited()
        self.assertIn("8MB", interaction.edit_original_response.call_args.kwargs["content"])

    async def test_library_buttons_check_owner(self):
        interaction = SimpleNamespace(user=SimpleNamespace(id=2), response=SimpleNamespace(send_message=AsyncMock()))
        view = LibraryView("guild", "1", [])
        self.assertFalse(await view.interaction_check(interaction))
        self.assertTrue(interaction.response.send_message.call_args.kwargs["ephemeral"])

    async def test_modal_inheritance_and_component_payloads(self):
        message = SimpleNamespace(content="hello", attachments=[])
        modal = MessageQuestionModal("owner", message, None)
        self.assertEqual(len(modal.children), 1)
        self.assertTrue(modal.to_components())
        self.assertTrue(SnoozeModal(1).to_components())
        self.assertTrue(ReminderView(1).to_components())
        view = LibraryView("guild", "owner", [])
        view.build_embed()
        self.assertTrue(all(item.disabled for item in view.children))

    async def test_source_question_uses_bounded_context_and_no_tools(self):
        source = "intro " * 6000 + "DOCKER_SOLUTION: recreate the container."
        response = AsyncMock(return_value={"message": {"content": "Recreate it."}})
        with patch("clients.source_client.ollama_request", new=response):
            answer = await source_response(source, "Where is the Docker solution?", b"image")
        self.assertEqual(answer, "Recreate it.")
        messages = response.call_args.args[0]
        self.assertIn("DOCKER_SOLUTION", messages[1]["content"])
        self.assertLess(len(messages[1]["content"]), 17_000)
        self.assertIn("images", messages[1])
        self.assertNotIn("tools", response.call_args.kwargs)

class HelperTests(unittest.TestCase):
    def test_long_unicode_answers_keep_full_download_under_discord_limit(self):
        answer = "\U0001f680" * 2500
        preview, files = answer_parts(answer)
        self.assertLessEqual(len(preview.encode("utf-16-le")) // 2, 2000)
        self.assertEqual(len(files), 1)
        self.assertEqual(files[0].fp.read().decode("utf-8"), answer)
        files[0].close()
        self.assertEqual(shorten_text("hello"), "hello")
        self.assertTrue(answer_parts("  ")[0])

    def test_repeating_reminders_skip_missed_occurrences(self):
        self.assertIsNone(next_reminder_due_at(100, "none", now=1_000_000))
        daily = next_reminder_due_at(100, "daily", now=10 * 86400)
        self.assertGreater(daily, 10 * 86400)
        self.assertLessEqual(daily, 11 * 86400)
        weekly = next_reminder_due_at(100, "weekly", now=20 * 86400)
        self.assertGreater(weekly, 20 * 86400)
        self.assertLessEqual(weekly, 27 * 86400)

    def test_powerpoint_extraction_includes_slide_number(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Example presentation"
        data = BytesIO()
        presentation.save(data)
        result = extract_text("example.pptx", data.getvalue())
        self.assertIn("Slide 1:\n", result)
        self.assertIn("Example presentation", result)

    def test_source_excerpt_preserves_short_documents(self):
        self.assertEqual(select_source_text("Page 1: hello", "what is this?"), "Page 1: hello")

if __name__ == "__main__":
    unittest.main()
